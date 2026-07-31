"""
Ekstraksi teks dari materi kuliah (PPT/PDF) yang diupload dosen/mahasiswa,
supaya rangkuman LLM bisa "nyambung" ke istilah dan struktur slide asli
(bukan cuma dari hasil ASR yang bisa salah tangkap istilah teknis).
"""

from pathlib import Path


def extract_pptx_text(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    lines = []
    for i, slide in enumerate(prs.slides, start=1):
        slide_lines = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                slide_lines.append(shape.text_frame.text.strip())
        if slide_lines:
            lines.append(f"[Slide {i}] " + " | ".join(slide_lines))
    return "\n".join(lines)


def extract_pdf_text(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(f"[Hal. {i}] {text}")
    return "\n".join(pages)


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    try:
        if suffix == ".pptx":
            return extract_pptx_text(path)
        if suffix == ".pdf":
            return extract_pdf_text(path)
    except Exception as e:
        return f"[Gagal ekstrak {path.name}: {e}]"
    return ""


def extract_all(paths: list[Path], max_chars_per_file: int = 6000) -> str:
    chunks = []
    for p in paths:
        text = extract_text(p)
        if text:
            chunks.append(f"=== {p.name} ===\n{text[:max_chars_per_file]}")
    return "\n\n".join(chunks)
